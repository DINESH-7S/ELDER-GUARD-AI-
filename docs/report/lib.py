"""Small python-docx toolkit that reproduces the SNS mini-project report layout."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Times New Roman'
TEAL = RGBColor(0x0E, 0x47, 0x43)
FIG = r'E:\Elder_guard\docs\report\figures'
ASSETS = r'E:\Elder_guard\docs\presentation\assets'

doc = Document()
_tab_no = {}
_fig_no = {}


def _font(style, size, bold=False, italic=False, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)


def setup():
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
    st = doc.styles
    n = st['Normal']
    _font(n, 12)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.5
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, before, after in (('Heading 1', 16, 0, 14), ('Heading 2', 14, 14, 6), ('Heading 3', 12, 10, 4)):
        h = st[name]
        _font(h, size, bold=True, color=RGBColor(0, 0, 0))
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.line_spacing = 1.15
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ('TableCaption', 'FigureCaption'):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st['Normal']
        _font(s, 11, bold=True)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.space_after = Pt(10)
        s.paragraph_format.line_spacing = 1.15
    st['TableCaption'].paragraph_format.keep_with_next = True
    st['TableCaption'].paragraph_format.space_before = Pt(8)
    st['FigureCaption'].paragraph_format.space_before = Pt(2)
    cl = st.add_style('ChapterLabel', WD_STYLE_TYPE.PARAGRAPH)
    cl.base_style = st['Normal']
    _font(cl, 14, bold=True)
    cl.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cl.paragraph_format.space_before = Pt(60)
    cl.paragraph_format.space_after = Pt(4)
    cl.paragraph_format.keep_with_next = True
    cl.paragraph_format.page_break_before = True
    for name in ('List Bullet',):
        _font(st[name], 12)
        st[name].paragraph_format.line_spacing = 1.3
        st[name].paragraph_format.space_after = Pt(3)
        st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


setup()


def add_field(par, instr, placeholder=''):
    run = par.add_run()
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, placeholder), ('end', None)):
        if kind:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), kind)
            run._r.append(e)
        elif text is instr:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = instr
            run._r.append(e)
        else:
            t = OxmlElement('w:t')
            t.text = placeholder
            run._r.append(t)
    return run


def P(text='', bold=False, italic=False, align=None, size=None, after=None, before=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        _runs(p, text, bold, italic, size)
    if align == 'c':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'l':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif align == 'r':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def _runs(p, text, bold=False, italic=False, size=None):
    """Supports **bold** inline markup."""
    parts = text.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)


def B(items):
    for t in items:
        p = doc.add_paragraph(style='List Bullet')
        _runs(p, t)


def chapter(n, title):
    P(f'Chapter {n}', style='ChapterLabel')
    doc.add_heading(f'{n}. {title}', 1)
    global _cur
    _cur = n


def H2(t):
    doc.add_heading(t, 2)


def H3(t):
    doc.add_heading(t, 3)


_cur = 0


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def table(caption, headers, rows, widths=None, size=10.5):
    key = _cur
    _tab_no[key] = _tab_no.get(key, 0) + 1
    num = f'{key}.{_tab_no[key]}' if key else f'{_tab_no[key]}'
    P(f'Table {num}: {caption}', style='TableCaption')
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(2)
        _shade(c, '14625C')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            _runs(p, str(v), size=size)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.space_after = Pt(2)
    # repeat header row, avoid splitting rows
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for r in t.rows:
        pr = r._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        cs.set(qn('w:val'), 'true')
        pr.append(cs)
    if widths:
        t.autofit = False
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    P('', after=6)


def figure(path, caption, width_cm=14.5, prefix=None):
    key = prefix or _cur
    _fig_no[key] = _fig_no.get(key, 0) + 1
    num = f'{key}.{_fig_no[key]}'
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width_cm))
    P(f'Figure {num}: {caption}', style='FigureCaption')


def code(text):
    for line in text.strip('\n').split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(line if line else ' ')
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Consolas')
    P('', after=4)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_section(num_format, start=None, footer=True):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    pg = sec._sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sec._sectPr.append(pg)
    pg.set(qn('w:fmt'), num_format)
    if start is not None:
        pg.set(qn('w:start'), str(start))
    return sec


def footer_number(sec):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, 'PAGE', '1')
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(11)


def toc(instr):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_field(p, instr, 'Right-click and choose Update Field to build this list.')
    return p
