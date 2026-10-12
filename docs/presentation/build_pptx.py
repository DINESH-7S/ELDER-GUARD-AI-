"""Build the editable ElderGuard mini-project deck (.pptx).

  python docs/presentation/build_pptx.py

Everything is real, editable PowerPoint content: text boxes, rounded rectangles, circles.
Only the college logos, the app screenshots and the small icons are pictures.
Layout mirrors deck.template.html (1440 x 810 px design, 1 px = 0.75 pt).
"""
import json
import re
import tempfile
from pathlib import Path

import pymupdf as fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
TMP = Path(tempfile.mkdtemp(prefix="eg_deck_"))
ICONS = json.loads((ASSETS / "icons.json").read_text(encoding="utf8"))

# palette sampled from the original Team 8 template
CREAM, BROWN, INK, CARD, CARD_LINE, DARK, GREY = "FFFBFA", "8B5D52", "3F342E", "E9E2DE", "D8CFCA", "2A2420", "F2F2F2"
WHITE, SOFT = "FFFFFF", "F0DED7"
SERIF = "Times New Roman"

PX = 9525


def px(v):
    return Emu(int(round(v * PX)))


def rgb(hex_):
    return RGBColor.from_string(hex_)


def pt(px_size):
    return Pt(px_size * 0.75)


prs = Presentation()
prs.slide_width, prs.slide_height = px(1440), px(810)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ primitives
def new_slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def shape(slide, kind, x, y, w, h, fill=None, line=None, radius=None, name=None):
    sh = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def card(slide, x, y, w, h, brown=False, name="Card"):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=BROWN if brown else CARD,
                 line=None if brown else CARD_LINE, radius=14, name=name)


def circle(slide, x, y, d, fill=BROWN):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill)


def _bullet(paragraph, indent_px=26, char="•"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent_px * PX)))
    pPr.set("indent", str(-int(indent_px * PX)))
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu_font = etree.SubElement(pPr, qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def text(slide, x, y, w, h, paras, size=22, color=INK, align="l", anchor="t", font=SERIF,
         bold=False, italic=False, spacing=1.12, after=0, name=None):
    """paras: list of paragraphs. A paragraph is a string, or a list of (text, opts) runs,
    or ('bullet', <string-or-runs>). opts keys: color, bold, italic, size."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = spacing
        if after:
            p.space_after = Pt(after)
        bullet = False
        if isinstance(para, tuple) and para[0] == "bullet":
            bullet, para = True, para[1]
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
        if bullet:
            _bullet(p)
    return tb


def title(slide, lead, rest, one_line=False):
    runs = []
    if lead:
        runs.append((lead, {"color": BROWN}))
    runs.append((rest, {}))
    return text(slide, 77, 106, 1140, 110, [runs], size=42, spacing=1.0, name="Title")


def picture(slide, path, x, y, w, h, name=None):
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    if name:
        pic.name = name
    return pic


def cover_crop(src, w, h, anchor=(0.0, 0.0)):
    """Crop an image to the w:h aspect (object-fit: cover), anchored top-left by default."""
    im = Image.open(src).convert("RGB")
    sw, sh = im.size
    target = w / h
    if sw / sh > target:
        nw = int(sh * target)
        left = int((sw - nw) * anchor[0])
        im = im.crop((left, 0, left + nw, sh))
    else:
        nh = int(sw / target)
        top = int((sh - nh) * anchor[1])
        im = im.crop((0, top, sw, top + nh))
    out = TMP / (Path(src).stem + f"_{int(w)}x{int(h)}.png")
    im.save(out)
    return out


_icon_cache = {}


def icon_png(name, color=WHITE, size=256):
    key = (name, color)
    if key in _icon_cache:
        return _icon_cache[key]
    svg = ICONS[name].replace("currentColor", f"#{color}")
    svg = re.sub(r'width="1em"', f'width="{size}"', svg)
    svg = re.sub(r'height="1em"', f'height="{size}"', svg)
    doc = fitz.open(stream=svg.encode("utf8"), filetype="svg")
    pix = doc[0].get_pixmap(alpha=True)
    out = TMP / f"icon_{name}_{color}.png"
    pix.save(str(out))
    _icon_cache[key] = out
    return out


def icon_badge(slide, name, x, y, d=62, fill=BROWN):
    circle(slide, x, y, d, fill)
    s = d * 0.48
    picture(slide, icon_png(name), x + (d - s) / 2, y + (d - s) / 2, s, s, name=f"Icon {name}")


def logos(slide, grey=False):
    sfx = "-grey" if grey else ""
    picture(slide, ASSETS / f"crest{sfx}.png", 0, 0, 112, 112, name="College crest")
    picture(slide, ASSETS / f"badges{sfx}.png", 1222, 0, 218, 108, name="SNS and Design Thinking badges")


def footer(slide, txt, y=752, italic=False, size=22, x=58, w=1324):
    text(slide, x, y, w, 50, [txt], size=size, italic=italic, color=BROWN if italic else INK, name="Footer")


# ------------------------------------------------------------------ slides
# 1. Title
s = new_slide(); logos(s)
text(s, 0, 70, 1440, 70, ["SNS COLLEGE OF TECHNOLOGY"], size=46, align="c", name="College")
text(s, 0, 160, 1440, 40, ["23ITP202 MINI PROJECT- I"], size=28, align="c", name="Course")
card(s, 53, 296, 556, 404, brown=True, name="Tagline card")
picture(s, icon_png("Heartbeat"), 331 - 80, 340, 160, 160, name="Icon heartbeat")
text(s, 53, 560, 556, 90, ["Care that reaches them,", "wherever you are"], size=27, color=WHITE, align="c", spacing=1.35, name="Tagline")
text(s, 648, 296, 740, 110, ["ELDERGUARD AI: CONNECTED ELDER CARE PLATFORM"], size=40, spacing=1.05, name="Project name")
text(s, 648, 418, 760, 36, ["Problem Statement: Elder Safety, Health Monitoring and Connected Care"], size=22, italic=True, color=BROWN, name="Subtitle")
text(s, 745, 488, 500, 44, ["Presented By:"], size=30, name="Presented by")
text(s, 778, 556, 640, 200, ["713525AD015 - DEEPIKA S", "713525AD036 - HANA SHAJIMON", "713525AD041 - HARINI M", "713525AD052 - JANANI D"], size=30, spacing=1.2, name="Team")

# 2. Problem statement
s = new_slide(); logos(s)
text(s, 75, 158, 660, 26, ["PROBLEM STATEMENT · ELDER SAFETY"], size=17, bold=True, color=BROWN, name="Eyebrow")
text(s, 75, 196, 680, 130, ["Keeping Elders Safe:", "A Connected-Care Approach"], size=46, spacing=1.05, name="Title")
text(s, 75, 352, 660, 300,
     ["Many older adults live apart from their families. Missed medicines, unnoticed changes in health and slow emergency response put them at risk, while family, doctors, pharmacies and hospitals each work from separate information. We built one connected platform where everyone in an elder's care circle sees the same live record and can act in time."],
     size=24, spacing=1.25, name="Body")
card(s, 798, 176, 596, 232)
text(s, 826, 200, 540, 24, ["WHAT HAPPENS TODAY"], size=17, bold=True, color=BROWN, name="Label")
text(s, 826, 238, 540, 150, ["The 8:00 AM dose is missed. Blood pressure is not logged for days. A fall at night and no one knows."], size=25, italic=True, spacing=1.15, name="Today")
card(s, 798, 476, 596, 262, brown=True)
text(s, 826, 500, 540, 24, ["WHAT ELDERGUARD DOES"], size=17, bold=True, color=WHITE, name="Label")
text(s, 826, 538, 540, 190, ["A reminder rings until it is answered, and the caregiver sees the dose logged live. One tap on SOS alerts family first, then a second contact, the hospital desk and emergency services if nobody responds."], size=25, color=WHITE, spacing=1.15, name="ElderGuard does")
circle(s, 1068, 420, 56, BROWN).line.color.rgb = rgb(CREAM)
picture(s, icon_png("ArrowDown"), 1068 + 14, 420 + 14, 28, 28, name="Icon arrow")

# 3. Empathy
s = new_slide(); logos(s)
title(s, "Understanding the Empathy", " Phase: Connecting with Elders and Caregivers")
text(s, 77, 238, 380, 220, [[("Deep Elder & Caregiver Research", {"bold": True})],
     "Start with genuine curiosity about the lives of older adults and the family members who worry about them. Understand daily routines, fears of being alone and how comfortable they are with technology."], size=22, spacing=1.12, name="Research")
text(s, 502, 238, 380, 220, [[("Building Emotional Connections", {"bold": True})],
     "Go beyond demographics. What makes an elder hesitate to ask for help? What keeps a caregiver awake at night? What would make both feel in control?"], size=22, spacing=1.12, name="Connections")
text(s, 77, 534, 390, 150, [("bullet", "Interview elders and family caregivers"), ("bullet", "Observe medicine routines at home"), ("bullet", "Build empathy maps of feelings and needs")], size=21, after=3, name="Methods 1")
text(s, 502, 534, 380, 150, [("bullet", "Medicine-routine diary studies"), ("bullet", "Contextual enquiry in homes and care centres"), ("bullet", "Journey mapping of an emergency")], size=21, after=3, name="Methods 2")
for (lab, body, x, y) in [("SAYS", "\"Did she take her tablets? Is she okay?\"", 955, 232), ("THINKS", "\"What if something happens and I am not there?\"", 1190, 232),
                          ("FEELS", "Anxious, guilty, always on call", 955, 476), ("DOES", "Calls repeatedly, asks neighbours to check in", 1190, 476)]:
    card(s, x, y, 227, 232, name=f"Empathy {lab}")
    text(s, x + 22, y + 24, 185, 24, [lab], size=18, bold=True, color=BROWN, name=f"{lab} label")
    text(s, x + 22, y + 62, 185, 150, [body], size=22, name=f"{lab} text")
footer(s, "Empathy map of a family caregiver. Illustrative persona; interviews with real elders and caregivers are planned.", y=756, italic=True, size=20, x=77)

# 4. Define
s = new_slide(); logos(s)
title(s, "Define", " the Problem: Identifying Core Pain Points in Elder Care")
items = [("Pill", "Missed Medicines", "Several daily medicines are easy to forget, and the family finds out late, if at all."),
         ("Siren", "Slow Emergency Response", "A fall or sudden illness can go unnoticed. Help depends on someone happening to call."),
         ("Link", "Disconnected Information", "Family, doctor, pharmacy and hospital each hold a piece, so prescriptions, vitals and history are not shared."),
         ("DeviceMobile", "Hard-to-Use Technology", "Small text, crowded screens and complex steps leave out the people who need help most.")]
for i, (ic, h, b) in enumerate(items):
    x, y = 58 + (i % 2) * 674, 214 + (i // 2) * 194
    card(s, x, y, 650, 170, name=h)
    icon_badge(s, ic, x + 28, y + 28)
    text(s, x + 112, y + 24, 510, 44, [h], size=31, name=f"{h} title")
    text(s, x + 112, y + 72, 510, 90, [b], size=22, spacing=1.12, name=f"{h} text")
card(s, 58, 662, 1324, 112, brown=True, name="Problem statement")
text(s, 88, 678, 1264, 84, ["The problem statement emerges from synthesis: \"Elders and their families need a simple, connected way to take medicines on time, share health information and get help fast, so older adults can live independently with confidence.\""], size=24, color=WHITE, spacing=1.15, anchor="m", name="Synthesis")

# 5. Ideate
s = new_slide(); logos(s)
title(s, "Ideate", " Creative Solutions: Brainstorming Elder Care Features")
ideas = [("Pill", "Reminders That Get Answered", "One gentle alarm per dose time. A tap logs the dose, and repeated misses alert the caregiver."),
         ("Siren", "One-Tap SOS with Escalation", "Family first, then a second contact, the hospital desk and emergency services on a timer."),
         ("Robot", "Vitals and AI Companion", "Track blood pressure, sugar, oxygen and heart rate, and chat with Mevi, a patient AI companion.")]
for i, (ic, h, b) in enumerate(ideas):
    x = 58 + i * 449.3
    card(s, x, 214, 425.3, 270, name=h)
    icon_badge(s, ic, x + 28, 240)
    text(s, x + 28, 318, 372, 40, [h], size=28, name=f"{h} title")
    text(s, x + 28, 364, 372, 110, [b], size=21, spacing=1.12, name=f"{h} text")
text(s, 58, 548, 800, 36, ["More ideas from the brainstorm"], size=28, color=BROWN, name="More ideas heading")
text(s, 58, 596, 640, 100, [("bullet", "Prescription scanner that turns a photo into a medicine schedule"), ("bullet", "Doctor appointments with queue tokens")], size=22, after=4, name="More ideas 1")
text(s, 720, 596, 660, 100, [("bullet", "Doctor-to-pharmacy e-prescriptions"), ("bullet", "Five role-based portals with a large-text mode")], size=22, after=4, name="More ideas 2")
footer(s, "Innovation thrives when we encourage wild ideas, defer judgement and build on each other's concepts. Generate quantity first, then refine for safety and quality.", y=754, italic=True, size=20)

# 6. Proposed solution
s = new_slide(); logos(s)
title(s, "Proposed Solution:", " How ElderGuard Works End to End")
steps = [("Pill", "1. Capture", "Elders log medicines and vitals; doctors write prescriptions"),
         ("Link", "2. Connect", "One shared record links elder, caregiver, doctor, pharmacy and hospital"),
         ("Bell", "3. Alert", "Reminders ring, unusual readings are flagged, SOS goes out in one tap"),
         ("Siren", "4. Escalate", "Caregiver now, second contact at 2 min, hospital at 5, emergency services at 10"),
         ("ArrowsClockwise", "5. Review", "History and records let family and doctors act on what happened")]
for i, (ic, h, b) in enumerate(steps):
    x = 58 + i * 274
    card(s, x, 214, 232, 308, name=h)
    icon_badge(s, ic, x + 85, 238)
    text(s, x + 10, 316, 212, 40, [h], size=30, align="c", name=f"{h} title")
    text(s, x + 18, 366, 196, 140, [b], size=19, align="c", spacing=1.12, name=f"{h} text")
    if i < 4:
        picture(s, icon_png("CaretRight", BROWN), x + 232 + 8, 214 + 140, 26, 26, name="Arrow")
card(s, 58, 548, 1324, 190, brown=True, name="Safety guardrails")
text(s, 92, 574, 1200, 40, [[("Safety guardrails built in", {"bold": True})]], size=30, color=WHITE, name="Guardrails title")
text(s, 92, 626, 1250, 100, [("bullet", "Role-based access: each user sees only the elders they are linked to"), ("bullet", "ElderGuard alerts and informs; it does not diagnose or replace a doctor or emergency services")], size=22, color=WHITE, after=6, name="Guardrails")

# 7. Prototype
s = new_slide(); logos(s)
title(s, "Prototype", " Development: Building Tangible Solutions for User Testing")
card_d = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 58, 214, 610, 520, fill=DARK, radius=14, name="Prototype frame")
text(s, 80, 236, 560, 28, [[("ElderGuard AI · working prototype", {"bold": True})]], size=20, color="E7D3CB", name="Frame caption")
picture(s, cover_crop(ASSETS / "shot-elder.jpg", 566, 312), 80, 272, 566, 312, name="Elder portal screenshot")
tw = (566 - 28) / 3
for i, (f, cap) in enumerate([("shot-caregiver.jpg", "Caregiver portal"), ("shot-pharmacy.jpg", "Pharmacy queue"), ("shot-landing.jpg", "Landing page")]):
    x = 80 + i * (tw + 14)
    picture(s, cover_crop(ASSETS / f, tw, 100), x, 598, tw, 100, name=cap)
    text(s, x, 702, tw, 22, [cap], size=16, color="E7D3CB", align="c", name=f"{cap} caption")
protos = [("PencilRuler", "Design System", "Warm paper background, one teal accent, large readable type and a text-size control for low vision."),
          ("Cursor", "Interactive Portals", "Five role-based portals (elder, caregiver, doctor, hospital, pharmacy) running on real data."),
          ("Cube", "Minimum Viable Product", "A working React and FastAPI app: reminders, SOS escalation, vitals, and prescriptions that reach the pharmacy.")]
for i, (ic, h, b) in enumerate(protos):
    y = 214 + i * 150
    icon_badge(s, ic, 712, y)
    text(s, 796, y - 4, 586, 44, [h], size=33, name=f"{h} title")
    text(s, 796, y + 44, 586, 90, [b], size=21, spacing=1.12, name=f"{h} text")
text(s, 712, 690, 670, 60, ["Prototypes are learning tools. Build fast, fail fast and iterate based on what real users show you."], size=21, name="Footer")

# 8. Testing
s = new_slide(); logos(s)
title(s, "Testing", " Phase: Verifying the Prototype and Planning User Trials")
text(s, 77, 198, 790, 44, ["Automated Testing"], size=34, name="Heading 1")
text(s, 77, 252, 790, 90, ["An automated browser opens every portal page, the demo story is checked end to end, and each screen is scanned against WCAG 2.1 accessibility rules."], size=22, spacing=1.15, name="Automated testing")
text(s, 77, 362, 790, 44, ["Key Testing Methods"], size=34, name="Heading 2")
text(s, 77, 420, 800, 220, [
    [("End-to-end tests (Playwright): ", {"bold": True}), ("22 tests across all five portals and phone size", {})],
    [("Accessibility audit (axe): ", {"bold": True}), ("0 serious violations against WCAG 2.1 A/AA", {})],
    [("SOS escalation drill: ", {"bold": True}), ("all 4 tiers fired in order in a timed run", {})],
    [("Planned: ", {"bold": True}), ("teach-back and think-aloud sessions with real elders and caregivers", {})]], size=22, spacing=1.15, after=6, name="Methods")
for (cx, cy, big, lab, ly) in [(960, 180, "22/22", "Automated tests passing", 276), (960, 466, "4/4", "SOS tiers fired in the timed drill", 562)]:
    circle(s, cx, cy, 270, BROWN)
    circle(s, cx + 52, cy + 52, 166, CREAM)
    text(s, cx + 52, cy + 52, 166, 166, [big], size=52, align="c", anchor="m", name=f"{big} value")
    text(s, 1236, ly, 190, 70, [lab], size=22, spacing=1.1, name=f"{big} label")
footer(s, "Testing is discovery, not validation. Automated checks show that it works; sessions with real elders will show whether it helps.", y=752)

# 9. Use case
s = new_slide(); logos(s)
title(s, "Use Case:", " How Connected Care Protects an Elder at Home")
shape(s, MSO_SHAPE.RECTANGLE, 722, 214, 2, 540, fill=CARD_LINE, name="Timeline")
left = [(214, "Initial Challenge", "Mrs. Anaya Rao, 77, lives alone and takes three medicines a day. Her daughter Maya lives in another city and worries about missed doses and emergencies."),
        (500, "Design Response", "A reminder rings and Anaya taps Taken; Maya sees it live. Dr. Mehta prescribes in the app and the order reaches the pharmacy. One tap on SOS alerts Maya first, then the hospital.")]
right = [(366, "Empathy Insights", "Anaya does not want to be a burden and Maya cannot call every hour. They want reassurance that does not feel like surveillance."),
         (640, "Expected Impact", "Fewer missed doses, faster response and less worry for the family. To be measured in a pilot.")]
for y, h, b in left:
    text(s, 70, y, 575, 40, [h], size=33, align="r", name=h)
    text(s, 70, y + 46, 575, 120, [b], size=21, align="r", spacing=1.12, name=f"{h} text")
for y, h, b in right:
    text(s, 786, y, 590, 40, [h], size=33, name=h)
    text(s, 786, y + 46, 590, 110, [b], size=21, spacing=1.12, name=f"{h} text")
for n, y in enumerate([214, 366, 500, 640], 1):
    r = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 692, y, 62, 62, fill=BROWN, radius=10, name=f"Step {n}")
    text(s, 692, y, 62, 62, [str(n)], size=30, color=WHITE, align="c", anchor="m", name=f"Step {n} number")
footer(s, "Illustrative scenario based on this project's demo data; impact is a target to be validated.", y=766, italic=True, size=19)

# 10. Integration challenges
s = new_slide(); logos(s)
title(s, "Integration", " Challenges: Balancing Elder Needs with Technical and Clinical Constraints")
chal = [("Lock", "Privacy & Compliance", "Health data is highly sensitive. It needs consent, encryption and access logs in line with India's DPDP Act. Role-based access is in place; consent and audit features come next."),
        ("Siren", "Emergency Reliability", "SOS must never stall. Alerts are in-app and push today; SMS and voice-call providers and a durable job queue are needed for production."),
        ("Database", "Data & Integration", "Records come from elders, doctors, pharmacies and hospitals. One shared data model and clean APIs keep them consistent."),
        ("WifiHigh", "Access & Performance", "Older users often have older phones and slow networks. Large text and simple screens are prioritised; an installable app is next.")]
for i, (ic, h, b) in enumerate(chal):
    x, y = 58 + (i % 2) * 451, 214 + (i // 2) * 258
    card(s, x, y, 429, 236, name=h)
    icon_badge(s, ic, x + 24, y + 22, 56)
    text(s, x + 96, y + 30, 320, 40, [h], size=30, name=f"{h} title")
    text(s, x + 24, y + 96, 385, 130, [b], size=20, spacing=1.12, name=f"{h} text")
card(s, 976, 214, 406, 494, brown=True, name="Safety first")
picture(s, icon_png("ShieldCheck"), 1179 - 70, 290, 140, 140, name="Icon shield")
text(s, 976, 458, 406, 50, ["Safety first"], size=38, color=WHITE, align="c", name="Safety first title")
text(s, 1010, 520, 338, 150, ["ElderGuard alerts and informs; it does not diagnose or replace a doctor or emergency services."], size=23, color=WHITE, align="c", spacing=1.15, name="Safety first text")
footer(s, "Successful integration needs continuous dialogue between clinicians, designers and engineers. Constraints often spark creative solutions when teams collaborate early.", y=724, size=21)

# 11. Success metrics
s = new_slide(); logos(s)
title(s, "", "Success Metrics: Measuring the Impact of ElderGuard")
mets = [("22/22", "Tests Passing", "Automated end-to-end checks across all five portals"),
        ("0", "Accessibility Issues", "Serious issues found by an axe audit (WCAG 2.1 A/AA) on every page"),
        ("4", "Escalation Tiers", "Caregiver, second contact, hospital, emergency services, verified in a timed drill"),
        ("90%+", "Doses On Time", "Pilot target for doses logged when due, to be validated")]
for i, (n, h, b) in enumerate(mets):
    x = 58 + i * 331
    text(s, x, 206, 331, 90, [n], size=70, color=BROWN, align="c", name=f"{h} value")
    text(s, x, 306, 331, 44, [h], size=32, align="c", name=f"{h} title")
    text(s, x + 22, 360, 287, 110, [b], size=21, align="c", spacing=1.12, name=f"{h} text")
text(s, 58, 478, 1324, 30, ["The first three are measured on the prototype; the last is a target for the pilot, to be validated with real users"], size=20, italic=True, color=BROWN, name="Note")
text(s, 77, 540, 900, 46, ["Qualitative Indicators of Success"], size=34, name="Qualitative heading")
text(s, 77, 598, 640, 140, [("bullet", "Caregivers report less worry and fewer check-in calls"), ("bullet", "Elders use the SOS button and large-text mode without help")], size=22, after=4, name="Qualitative 1")
text(s, 740, 598, 640, 140, [("bullet", "Doctors and pharmacies adopt the shared prescription flow"), ("bullet", "Shorter time from SOS to first response")], size=22, after=4, name="Qualitative 2")

# 12. Future enhancements
s = new_slide(); logos(s)
title(s, "Future Enhancements:", " Expanding ElderGuard's Reach and Reliability")
future = [("Phone", "Phase 1", "SMS & Voice-Call Alerts", "Reach caregivers and hospitals even without the app, with delivery receipts and retries."),
          ("SealCheck", "Phase 1", "Verified Doctors & Hospitals", "Licence checks before a clinical account can issue prescriptions."),
          ("Microphone", "Phase 2", "Multilingual Voice Assistant", "Speak and listen in Tamil, Hindi and other regional languages."),
          ("DeviceMobile", "Phase 2", "Installable Mobile App", "Push notifications, offline mode and an always-ready SOS button."),
          ("Watch", "Phase 3", "Wearable Integration", "Read heart rate and blood pressure, and detect falls, from smart watches."),
          ("Link", "Phase 3", "EHR & ABDM Integration", "Pull records securely from hospitals and ABHA health IDs, with consent.")]
for i, (ic, ph, h, b) in enumerate(future):
    x, y = 58 + (i % 3) * 448.7, 214 + (i // 3) * 254
    card(s, x, y, 426.7, 232, name=h)
    icon_badge(s, ic, x + 24, y + 22, 56)
    text(s, x + 426.7 - 24 - 120, y + 24, 120, 26, [ph], size=18, bold=True, color=BROWN, align="r", name=f"{h} phase")
    text(s, x + 24, y + 98, 380, 44, [h], size=30, name=f"{h} title")
    text(s, x + 24, y + 146, 380, 80, [b], size=20, spacing=1.12, name=f"{h} text")
footer(s, "Roadmap: Phase 1 reliable alerts and trusted accounts → Phase 2 voice and mobile → Phase 3 wearables and health-record integration.", y=752, italic=True, size=20)

# 13. Next steps
s = new_slide(); logos(s)
title(s, "Next Steps:", " Taking ElderGuard from Prototype to Pilot")
nxt = [("Users", "Validate with Elders and Families", "Run usability sessions with older adults and caregivers, and refine the screens from what they struggle with."),
       ("ShieldCheck", "Harden for Real Use", "Add SMS and call alerts, secure secrets and logins, privacy consent and monitoring."),
       ("Rocket", "Pilot and Scale", "Partner with a clinic or care home for a small pilot, measure results, then expand.")]
for i, (ic, h, b) in enumerate(nxt):
    y = 214 + i * 172
    card(s, 58, y, 880, 152, name=h)
    icon_badge(s, ic, 86, y + 45)
    text(s, 172, y + 26, 740, 44, [h], size=34, name=f"{h} title")
    text(s, 172, y + 76, 740, 70, [b], size=21, spacing=1.12, name=f"{h} text")
card(s, 976, 214, 406, 494, brown=True, name="In summary")
text(s, 1008, 248, 340, 44, ["In Summary"], size=34, color=WHITE, name="Summary title")
text(s, 1008, 320, 346, 230, ["Connected information", "→ faster response", "→ safer elders", "→ peace of mind"], size=29, color=WHITE, spacing=1.35, name="Summary flow")
text(s, 1008, 600, 340, 80, ["Start small, measure impact, and scale what works."], size=22, italic=True, color=SOFT, name="Summary note")
footer(s, "Elder care is an ongoing commitment to staying close to the people we serve and improving relentlessly.", y=752, size=21)

# 14. Thank you
s = new_slide(GREY); logos(s, grey=True)
text(s, 0, 322, 1440, 110, ["Thank you!"], size=81, bold=True, color="1D1D1D", font="Arial", align="c", name="Thank you")
text(s, 0, 448, 1440, 44, ["ElderGuard AI: Care that reaches them, wherever you are"], size=26, color="555555", font="Arial", align="c", name="Tagline")

out = HERE / "ElderGuard_Mini_Project_Presentation.pptx"
prs.save(out)
print("saved", out, "| slides:", len(prs.slides))
